import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import logging
import io
from typing import Dict, Any, Optional

logger = logging.getLogger(__name__)

class TemplateGenerator:
    """
    Handles the generation of IND document templates based on provided data
    """
    
    def __init__(self):
        self.templates_dir = os.path.join(os.path.dirname(__file__), "templates")
        self.output_dir = os.path.join(os.path.dirname(__file__), "output")
        
        # Ensure output directory exists
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)
            
        logger.info(f"Initialized TemplateGenerator with templates dir: {self.templates_dir}")
    
    def generate_module3(self, data: Dict[str, Any]) -> Optional[bytes]:
        """
        Generate Module 3 (Chemistry, Manufacturing, and Controls) document
        
        Args:
            data: Dictionary containing CMC data for Module 3
            
        Returns:
            Document as bytes or None if generation fails
        """
        try:
            drug_name = data.get('drug_name', 'Unknown')
            logger.info(f"Generating Module 3 document for drug: {drug_name}")
            
            # Creating a blank document to ensure proper formatting
            doc = Document()
            
            # Set up document properties
            doc.core_properties.title = f"Module 3 - CMC - {drug_name}"
            doc.core_properties.subject = "Chemistry, Manufacturing, and Controls"
            
            # Add title
            title = doc.add_heading("MODULE 3: CHEMISTRY, MANUFACTURING, AND CONTROLS", level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Add drug name
            drug_name_paragraph = doc.add_paragraph()
            drug_name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            drug_name_run = drug_name_paragraph.add_run(drug_name)
            drug_name_run.font.size = Pt(18)
            drug_name_run.bold = True
            
            # Add IND application subtitle
            subtitle = doc.add_paragraph("Investigational New Drug Application")
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Add batch info
            batch_para = doc.add_paragraph()
            batch_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            batch_para.add_run(f"Batch Number: {data.get('batch_number', 'N/A')}")
            
            doc.add_paragraph()  # Add spacing
            
            # Section 1: Drug Substance
            doc.add_heading("1. DRUG SUBSTANCE", level=1)
            
            # Manufacturing Process
            doc.add_heading("1.1 Manufacturing Information", level=2)
            doc.add_paragraph(f"Manufacturing Site: {data.get('manufacturing_site', 'N/A')}")
            
            # Section 2: Specifications and Analytical Methods
            doc.add_heading("2. SPECIFICATIONS AND ANALYTICAL METHODS", level=1)
            
            # Create specifications table
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            
            # Add header row
            header_cells = table.rows[0].cells
            header_cells[0].text = "Parameter"
            header_cells[1].text = "Limit"
            header_cells[2].text = "Results"
            
            # Make header row bold
            for cell in header_cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Add specification rows
            for spec in data.get('specifications', []):
                row_cells = table.add_row().cells
                row_cells[0].text = spec.get('parameter', 'N/A')
                row_cells[1].text = spec.get('limit', 'N/A')
                row_cells[2].text = spec.get('result', 'N/A')
                
                # Center content
                for i in range(1, 3):  # Center the last two columns
                    for paragraph in row_cells[i].paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Section 3: Stability Data
            doc.add_heading("3. STABILITY DATA", level=1)
            
            # Create stability table
            stab_table = doc.add_table(rows=1, cols=2)
            stab_table.style = 'Table Grid'
            
            # Add header row
            stab_header = stab_table.rows[0].cells
            stab_header[0].text = "Time Point"
            stab_header[1].text = "Result"
            
            # Make header row bold
            for cell in stab_header:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Add stability data rows
            for stab_data in data.get('stability_data', []):
                row_cells = stab_table.add_row().cells
                row_cells[0].text = stab_data.get('timepoint', 'N/A')
                row_cells[1].text = stab_data.get('result', 'N/A')
                
                # Center all cells
                for cell in row_cells:
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Save document to bytes
            output = io.BytesIO()
            doc.save(output)
            output.seek(0)
            
            # Also save a copy to the output directory for reference
            filename = f"Module3_CMC_{drug_name.replace(' ', '_')}.docx"
            filepath = os.path.join(self.output_dir, filename)
            doc.save(filepath)
            logger.info(f"Successfully saved Module 3 document to {filepath}")
            
            return output.getvalue()
            
        except Exception as e:
            logger.error(f"Error generating Module 3 document: {str(e)}", exc_info=True)
            return None
            
    def save_module3_to_file(self, data: Dict[str, Any]) -> Optional[str]:
        """
        Generate Module 3 document and save it to a file
        
        Args:
            data: Dictionary containing CMC data
            
        Returns:
            Path to the saved file or None if generation fails
        """
        try:
            doc_bytes = self.generate_module3(data)
            if not doc_bytes:
                return None
                
            drug_name = data.get('drug_name', 'Unknown')
            filename = f"Module3_CMC_{drug_name.replace(' ', '_')}.docx"
            filepath = os.path.join(self.output_dir, filename)
            
            with open(filepath, 'wb') as f:
                f.write(doc_bytes)
                
            logger.info(f"Successfully saved Module 3 document to {filepath}")
            return filepath
        except Exception as e:
            logger.error(f"Error saving Module 3 document: {str(e)}", exc_info=True)
            return None
            
    def batch_generate(self, batch_data: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
        """
        Generate multiple Module 3 documents in batch mode
        
        Args:
            batch_data: Dictionary mapping project IDs to their data
            
        Returns:
            Dictionary with status for each project ID
        """
        results = {}
        
        for project_id, data in batch_data.items():
            try:
                filepath = self.save_module3_to_file(data)
                if filepath:
                    results[project_id] = {
                        "status": "success",
                        "filepath": filepath,
                        "filename": os.path.basename(filepath)
                    }
                else:
                    results[project_id] = {
                        "status": "error",
                        "message": "Failed to generate document"
                    }
            except Exception as e:
                results[project_id] = {
                    "status": "error",
                    "message": str(e)
                }
                
        return results