from pathlib import Path
import argparse, re, html, textwrap, shutil, sys
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Preformatted
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
parser = argparse.ArgumentParser(description='Render a repository study without replacing its historical exports.')
parser.add_argument('source', nargs='?', default='docs/reports/DATA_ROOM_AI_MOAT_STUDY_2026-10-14.md', help='Markdown source relative to the repository root, or an absolute path inside it.')
parser.add_argument('stem', nargs='?', default='Concept2Cure-data-room-study-2026-10-14', help='A filename stem, never a path.')
parser.add_argument('--font-dir', type=Path, default=Path('/usr/share/fonts/truetype/dejavu'), help='Directory holding DejaVuSans.ttf, DejaVuSans-Bold.ttf and DejaVuSansMono.ttf.')
parser.add_argument('--check', action='store_true', help='Validate source, filename, fonts and document parsing without writing exports.')
args = parser.parse_args()
if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._-]{0,199}', args.stem):
    parser.error('stem must be a plain filename of at most 200 characters, without path separators')
SRC = Path(args.source)
SRC = (SRC if SRC.is_absolute() else ROOT / SRC).resolve()
try:
    SRC.relative_to(ROOT)
except ValueError:
    parser.error('source must be inside the repository')
if not SRC.is_file():
    parser.error('source must identify an existing Markdown file')
OUT = Path(__file__).resolve().parent / 'exports' / 'rebuilt'
STEM = args.stem
text = SRC.read_text(encoding='utf-8')
if not text.strip():
    parser.error('source must not be empty')
TITLE = text.splitlines()[0].removeprefix('# ')
fontdir = args.font_dir.expanduser().resolve()
for name, filename in [('StudySans','DejaVuSans.ttf'),('StudySansBold','DejaVuSans-Bold.ttf'),('StudyMono','DejaVuSansMono.ttf')]:
    if not (fontdir / filename).is_file():
        parser.error('required DejaVu font missing; supply --font-dir with the three font files')
    pdfmetrics.registerFont(TTFont(name, str(fontdir/filename)))
pdfmetrics.registerFontFamily('StudySans',normal='StudySans',bold='StudySansBold',italic='StudySans',boldItalic='StudySansBold')
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='StudyBody',fontName='StudySans',fontSize=9,leading=13,spaceAfter=6,textColor=colors.HexColor('#233447')))
styles.add(ParagraphStyle(name='StudyCell',fontName='StudySans',fontSize=7.6,leading=10.4,spaceAfter=0))
styles.add(ParagraphStyle(name='StudyTitle',fontName='StudySansBold',fontSize=21,leading=26,spaceAfter=14,textColor=colors.HexColor('#143c59')))
styles.add(ParagraphStyle(name='StudyH2',fontName='StudySansBold',fontSize=14,leading=18,spaceBefore=16,spaceAfter=8,keepWithNext=True,textColor=colors.HexColor('#143c59')))
styles.add(ParagraphStyle(name='StudyH3',fontName='StudySansBold',fontSize=11,leading=15,spaceBefore=11,spaceAfter=6,keepWithNext=True,textColor=colors.HexColor('#185c70')))
styles.add(ParagraphStyle(name='StudyCode',fontName='StudyMono',fontSize=7.8,leading=11,spaceAfter=10,backColor=colors.HexColor('#f2f5f8'),borderPadding=8))

def artifact_link(url):
    # Relative repository evidence links resolve inside the extracted bundle.
    if url.startswith(('https:', 'http:', 'mailto:', '#')):
        return url
    local, separator, anchor = url.partition('#')
    resolved = (SRC.parent / local).resolve()
    try:
        relative = resolved.relative_to(ROOT)
    except ValueError:
        return url
    return 'repository/' + relative.as_posix() + (separator + anchor if separator else '')

def inline(s):
    parts = []
    pos = 0
    for match in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',s):
        parts.append(html.escape(s[pos:match.start()]))
        parts.append('<link href="'+html.escape(artifact_link(match.group(2)),quote=True)+'" color="#146c8c">'+html.escape(match.group(1))+'</link>')
        pos=match.end()
    parts.append(html.escape(s[pos:]))
    value=''.join(parts)
    value=re.sub(r'\*\*([^*]+)\*\*',r'<b>\1</b>',value)
    return re.sub(r'`([^`]+)`',r'<font name="StudyMono">\1</font>',value)

def blocks(s):
    lines=s.splitlines(); i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip(): i+=1; continue
        if line.startswith('```'):
            code=[]; i+=1
            while i<len(lines) and not lines[i].startswith('```'): code.append(lines[i]); i+=1
            yield 'code','\n'.join(code); i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c or 'x') for c in row): rows.append(row)
                i+=1
            yield 'table',rows; continue
        if line.startswith('#'):
            m=re.match(r'^(#+)\s+(.*)',line)
            yield 'heading',(len(m.group(1)),m.group(2)); i+=1; continue
        if re.match(r'^(- |\d+\. )',line):
            yield 'list',line; i+=1; continue
        para=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|- |\d+\. )',lines[i]):
            para.append(lines[i]); i+=1
        yield 'para',' '.join(para)

parsed=list(blocks(text))
if args.check:
    print({'source': SRC.relative_to(ROOT).as_posix(), 'stem': STEM, 'blocks': len(parsed), 'writes': False})
    sys.exit(0)
OUT.mkdir(parents=True, exist_ok=True)
flow=[]
width=A4[0]-96
for kind,value in parsed:
    if kind=='heading':
        level,title=value
        flow.append(Paragraph(inline(title),styles['StudyTitle' if level==1 else 'StudyH2' if level==2 else 'StudyH3']))
    elif kind=='table':
        n=len(value[0])
        fractions={2:[.52,.48],3:[.22,.43,.35],4:[.16,.29,.32,.23]}.get(n,[1/n]*n)
        data=[[Paragraph(inline(c),styles['StudyCell']) for c in row] for row in value]
        table=LongTable(data,colWidths=[width*f for f in fractions],repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e0edf3')),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f7f9fb')]),
            ('LINEBELOW',(0,0),(-1,0),.7,colors.HexColor('#7496a8')),
            ('LINEBELOW',(0,1),(-1,-1),.2,colors.HexColor('#cad6de')),
            ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),
            ('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        flow.extend([table,Spacer(1,9)])
    elif kind=='code':
        wrapped='\n'.join('\n'.join(textwrap.wrap(x,92,replace_whitespace=False,drop_whitespace=False)) or ' ' for x in value.splitlines())
        flow.append(Preformatted(wrapped,styles['StudyCode']))
    else:
        flow.append(Paragraph(inline(value),styles['StudyBody']))

def page(canvas,doc):
    canvas.saveState()
    canvas.setFont('StudySans',7)
    canvas.setFillColor(colors.HexColor('#62798a'))
    canvas.drawString(48,27,'Concept2Cure | Evidence and filing enhancements | 14 October 2026')
    canvas.drawRightString(A4[0]-48,27,str(doc.page))
    canvas.restoreState()

pdf=OUT/(STEM+'.pdf')
SimpleDocTemplate(str(pdf),pagesize=A4,rightMargin=48,leftMargin=48,topMargin=44,bottomMargin=44,title=TITLE,author='Concept2Cure research').build(flow,onFirstPage=page,onLaterPages=page)

doc=Document()
section=doc.sections[0]
section.top_margin=Inches(.65); section.bottom_margin=Inches(.65)
section.left_margin=Inches(.7); section.right_margin=Inches(.7)
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(6)

def hyperlink(p,label,url):
    rel=p.part.relate_to(artifact_link(url),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'146C8C'); props.append(color); r.append(props)
    t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)

def add_inline(p,s):
    pos=0
    for m in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',s):
        p.add_run(s[pos:m.start()].replace('`','')); hyperlink(p,m.group(1),m.group(2)); pos=m.end()
    p.add_run(s[pos:].replace('`',''))

for kind,value in parsed:
    if kind=='heading':
        level,title=value; doc.add_heading(title,level=0 if level==1 else level-1)
    elif kind=='table':
        table=doc.add_table(rows=1,cols=len(value[0])); table.style='Light Shading Accent 1'
        for j,c in enumerate(value[0]): add_inline(table.rows[0].cells[j].paragraphs[0],c)
        for row in value[1:]:
            cells=table.add_row().cells
            for j,c in enumerate(row): add_inline(cells[j].paragraphs[0],c)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs: r.font.size=Pt(8)
        doc.add_paragraph()
    elif kind=='code':
        p=doc.add_paragraph(); r=p.add_run(value); r.font.name='Consolas'; r.font.size=Pt(8)
    else:
        p=doc.add_paragraph(); add_inline(p,value)
footer=section.footer.paragraphs[0]
footer.add_run('Concept2Cure | 14 October 2026')
doc.save(OUT/(STEM+'.docx'))

shutil.copyfile(SRC,OUT/(STEM+'.md'))
print({'pdf':str(pdf),'blocks':len(parsed),'bytes':pdf.stat().st_size})
